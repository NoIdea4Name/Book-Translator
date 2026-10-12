"""Export editable book text to Word; never include API keys or AI memory."""
from io import BytesIO
import os
from pathlib import Path
import tempfile


def export_docx(book, mode='translation', *, show_page_headers=True,
                show_placeholders=True, page_breaks=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    if mode not in ('source', 'translation', 'both'):
        raise ValueError('Unknown export format.')
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Normal', 'Heading 1', 'Heading 2'):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    emitted = False
    for index, page in enumerate(book.pages):
        source = page.source if page.source.strip() else ('[Not extracted]' if show_placeholders else '')
        translation = page.translation if page.translation.strip() else ('[Not translated]' if show_placeholders else '')
        blocks = [('SOURCE', source), (f'TRANSLATION ({page.target_language or "unspecified"})', translation)] if mode == 'both' else [('', source if mode == 'source' else translation)]
        if not show_page_headers and not any(text.strip() for _, text in blocks):
            continue
        if emitted and page_breaks:
            document.add_page_break()
        emitted = True
        if show_page_headers:
            heading = f'Page {index + 1}' + (f' · {page.name}' if page.name else '')
            if mode == 'translation':
                heading += f' ({page.target_language or "untranslated"})'
            if mode != 'source' and page.translation_is_stale:
                heading += ' [Source edited after translation]'
            document.add_heading(heading, level=1)
        for title, text in blocks:
            if title:
                document.add_heading(title, level=2)
            # Word paragraphs remain editable, including intentional blank lines.
            for line in text.split('\n'):
                document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def save_docx(path, data):
    """Replace only after a complete write; preserve an existing file on failure."""
    path = Path(path).resolve()
    temporary = None
    try:
        with tempfile.NamedTemporaryFile('wb', dir=path.parent, prefix=f'.{path.name}.', suffix='.tmp', delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
