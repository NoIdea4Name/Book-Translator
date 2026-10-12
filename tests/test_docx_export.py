from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from book_translator.core import BookDocument
from book_translator.docx_export import export_docx, save_docx


class DocxExportTests(unittest.TestCase):
    def setUp(self):
        path = Path('book.pdf')
        self.book = BookDocument([path], 'pdf', [(path, 0), (path, 1)])
        page = self.book.pages[0]
        page.name = 'Chương một'
        page.source = 'Original <text>\n\nSecond paragraph'
        page.translation = 'Xin chào thế giới\nDòng thứ hai'
        page.target_language = 'Vietnamese'
        page.translated_from = page.source

    def text(self, **options):
        return '\n'.join(p.text for p in Document(BytesIO(export_docx(self.book, **options))).paragraphs)

    def test_modes_headers_placeholders_unicode_and_no_memory(self):
        self.book.ai_summary = 'PRIVATE MEMORY'
        self.assertIn('Xin chào thế giới', self.text())
        self.assertIn('Page 1 · Chương một', self.text())
        self.assertIn('[Not translated]', self.text())
        plain = self.text(show_page_headers=False, show_placeholders=False)
        self.assertNotIn('Page 1', plain)
        self.assertNotIn('[Not translated]', plain)
        self.assertNotIn('PRIVATE MEMORY', plain)
        both = self.text(mode='both')
        self.assertIn('Original <text>\n\nSecond paragraph', both)
        self.assertIn('TRANSLATION (Vietnamese)', both)
        self.assertNotIn('Xin chào', self.text(mode='source'))

    def test_page_order_breaks_and_atomic_failure(self):
        self.book.pages[1].translation = 'Second page'
        self.book.reorder([1, 0])
        data = export_docx(self.book, page_breaks=True)
        doc = Document(BytesIO(data))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertLess(text.index('Second page'), text.index('Xin chào'))
        self.assertIn('w:type="page"', doc.element.xml)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'result.docx'
            path.write_bytes(b'old file')
            with patch('book_translator.docx_export.os.replace', side_effect=OSError('locked')):
                with self.assertRaises(OSError):
                    save_docx(path, data)
            self.assertEqual(path.read_bytes(), b'old file')
            self.assertEqual(len(list(Path(folder).iterdir())), 1)
